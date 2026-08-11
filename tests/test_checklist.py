import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from drawing_qa_agent.checklist import build_excel_checklist
from drawing_qa_agent.extractor import _extract_json_object, DrawingExtractorError
from drawing_qa_agent.schema import Dimension, DrawingAnalysis

SAMPLE_PATH = Path(__file__).parent / "sample_analysis.json"


@pytest.fixture()
def sample_analysis() -> DrawingAnalysis:
    data = json.loads(SAMPLE_PATH.read_text(encoding="utf-8"))
    return DrawingAnalysis.from_dict(data)


def test_dimension_limits_are_derived_when_missing():
    dim = Dimension(item_no=1, feature="測試", nominal_value=10.0, upper_tol=0.2, lower_tol=-0.1)
    assert dim.upper_limit == 10.2
    assert dim.lower_limit == 9.9


def test_build_excel_checklist_creates_expected_sheets_and_rows(tmp_path, sample_analysis):
    out_path = tmp_path / "checklist.xlsx"
    build_excel_checklist(
        sample_analysis, str(out_path), inspector="Antonio", inspection_date=date(2026, 8, 11)
    )

    assert out_path.exists()
    wb = load_workbook(out_path)
    assert wb.sheetnames == ["圖面基本資料", "尺寸檢核表"]

    info_ws = wb["圖面基本資料"]
    assert info_ws["A1"].value == "圖面判讀品保檢核表 — 基本資料"
    assert info_ws["B3"].value == sample_analysis.drawing_info.drawing_no

    checklist_ws = wb["尺寸檢核表"]
    header = [c.value for c in next(checklist_ws.iter_rows(min_row=1, max_row=1))]
    assert header[0] == "項次"
    assert header[15] == "判定結果"

    n = len(sample_analysis.dimensions)
    assert checklist_ws.max_row == n + 1
    # spot-check the first dimension row
    row2 = [c.value for c in next(checklist_ws.iter_rows(min_row=2, max_row=2))]
    assert row2[1] == sample_analysis.dimensions[0].feature
    assert row2[15].startswith("=IF(")


def test_extract_json_object_strips_markdown_fence():
    text = '```json\n{"drawing_info": {}, "dimensions": []}\n```'
    data = _extract_json_object(text)
    assert data == {"drawing_info": {}, "dimensions": []}


def test_extract_json_object_raises_on_garbage():
    with pytest.raises(DrawingExtractorError):
        _extract_json_object("not json at all")
